from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from codedir.models import CourseCatalogItem


CATALOG_FILENAME = "CE_Course Journal Project.xlsx"


def find_catalog_path(start: Path | None = None) -> Path | None:
    """Find the course catalog workbook near the app or current working folder."""
    roots: list[Path] = []
    if start:
        roots.append(start.resolve())
    roots.append(Path.cwd().resolve())
    roots.append(Path(__file__).resolve())

    for root in roots:
        for candidate_root in [root, *root.parents]:
            candidate = candidate_root / CATALOG_FILENAME
            if candidate.exists():
                return candidate
    return None


def load_course_catalog(catalog_path: Path) -> list[CourseCatalogItem]:
    workbook = load_workbook(catalog_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    if len(rows) < 3:
        return []

    semester_codes = [_clean(value) for value in rows[1][3:]]
    courses: list[CourseCatalogItem] = []

    for row in rows[2:]:
        course_number = _clean(row[0] if len(row) > 0 else "")
        course_name = _clean(row[1] if len(row) > 1 else "")
        assignment_spec = _clean_multiline(row[2] if len(row) > 2 else "")

        if not course_name and assignment_spec and courses:
            courses[-1].notes.append(assignment_spec)
            courses[-1].assignment_names = _assignment_names(
                courses[-1].assignment_spec,
                courses[-1].notes,
            )
            continue

        if not course_name:
            continue

        available_semesters: list[str] = []
        for code, marker in zip(semester_codes, row[3:]):
            if _clean(marker).casefold().startswith("x"):
                available_semesters.append(_semester_label(code))

        notes: list[str] = []
        item = CourseCatalogItem(
            course_number=course_number,
            course_name=course_name,
            assignment_spec=assignment_spec,
            available_semesters=available_semesters,
            assignment_names=[],
            notes=notes,
        )
        item.assignment_names = _assignment_names(assignment_spec, notes)
        courses.append(item)

    return courses


def _assignment_names(assignment_spec: str, notes: list[str]) -> list[str]:
    text = "\n".join(part for part in [assignment_spec, *notes] if part).strip()
    if not text:
        return []

    named_entries: list[tuple[int, str]] = []
    for line in text.splitlines():
        line = line.strip()
        match = re.search(r"\((\d{1,2})\)\s*$", line)
        if match:
            number = int(match.group(1))
            title = re.sub(r"\s*[-:]*\s*Journal Entry\s*\(\d{1,2}\)\s*$", "", line, flags=re.I)
            title = re.sub(r"\s*\(\d{1,2}\)\s*$", "", title).strip()
            named_entries.append((number, title or f"Assignment {number}"))

    if named_entries:
        return [f"Assignment {number} - {title}" for number, title in sorted(named_entries)]

    counts = _counted_assignment_groups(text)
    if not counts:
        return []

    names: list[str] = []
    for count, label in counts:
        base = _label_base(label)
        for index in range(1, count + 1):
            if len(counts) == 1:
                names.append(f"Assignment {index}")
            else:
                names.append(f"{base} {index}")
    return names


def _counted_assignment_groups(text: str) -> list[tuple[int, str]]:
    groups: list[tuple[int, str]] = []
    pattern = re.compile(
        r"(\d{1,2})\s*\"?\s*(COURSE JOURNAL|Journal Entry|CHANGEMAKER PROFILE|CASE STUDIES?)\"?",
        re.I,
    )
    for match in pattern.finditer(text):
        groups.append((int(match.group(1)), match.group(2)))
    return groups


def _label_base(label: str) -> str:
    value = label.casefold()
    if "changemaker" in value:
        return "Changemaker Profile"
    if "case" in value:
        return "Case Study"
    if "journal" in value:
        return "Assignment"
    return label.title()


def _semester_label(code: str) -> str:
    code = code.replace("\u202f", " ").strip()
    match = re.match(r"([A-Za-z]+)\s*(\d{2,4})", code)
    if not match:
        return code

    season_code = match.group(1).casefold()
    year = match.group(2)[-2:]
    season = {
        "f": "Fall",
        "s": "Spring",
        "ss": "Summer",
        "su": "Summer",
        "w": "Winter",
    }.get(season_code, season_code.title())
    return f"{season} {year}"


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\u202f", " ").replace("\xa0", " ").strip()


def _clean_multiline(value) -> str:
    value = _clean(value)
    return "\n".join(line.strip() for line in value.splitlines() if line.strip())
