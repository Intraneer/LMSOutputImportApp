from __future__ import annotations

import html
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from codedir.models import (
    ArchiveImportRequest,
    AssignmentMeta,
    CourseCatalogItem,
    ImportResult,
    WorkbookProfile,
    ZipImportPreview,
)


SEASON_ORDER = {
    "winter": 0,
    "spring": 1,
    "summer": 2,
    "fall": 3,
}

STOP_WORDS = {
    "a",
    "an",
    "and",
    "app",
    "assignment",
    "below",
    "box",
    "choose",
    "course",
    "define",
    "for",
    "from",
    "identify",
    "in",
    "is",
    "it",
    "journal",
    "networking",
    "of",
    "on",
    "or",
    "pick",
    "platform",
    "please",
    "question",
    "questions",
    "social",
    "step",
    "the",
    "to",
    "what",
    "with",
    "words",
    "your",
}

WORKSHEET_HEADERS = (
    "Semester",
    "Response",
    "Rubric_Version_Used",
    "Model_Version_Used",
    "AI_Overall_Level",
    "AI_Numeric_Score",
    "AI_Confidence",
    "AI_Reasoning",
    "AI_Feedback",
    "Flagged_For_Review",
)
WORKSHEET_HEADER_DESCRIPTORS = (
    "Existing column - untouched",
    "Existing column - untouched",
    "Ties this grade to the exact rubric version used",
    "For auditing drift over time",
    "One of this assignment's Achievement_Levels (Correct/Partial/Incorrect here)",
    "Mapped from point values in the rubric",
    "High/Low - Low auto-sets Flagged_For_Review",
    "Full JSON output from Step 3 grading prompt",
    "Student-facing feedback (2nd person)",
    "Auto-set Yes if AI_Confidence is Low",
)
METADATA_LABELS = (
    "Last Imported line",
    "Last Graded line",
    "Last Imported date",
    "Last Graded Date",
    "Last Imported Semester",
    "Last Graded Semester",
)
METADATA_NULL = "null"
ASSIGNMENT_NAME_ROW = 3
PROMPT_ROW = 4
RUBRIC_ROW = 5
MATCH_TITLE_COLUMN = 2
HEADER_ROW = 18
DESCRIPTOR_ROW = 19
FIRST_DATA_ROW = 20
ADVANCED_HISTORY_SHEET = "Advanced History"
GENERIC_ASSIGNMENT_PHRASES = (
    "Journal Entry",
    "Course Journal",
    "CHANGEMAKER PROFILE",
    "CASE STUDIES",
)


class _HTMLStripper(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._chunks: list[str] = []

    def handle_data(self, data: str) -> None:
        self._chunks.append(data)

    def get_text(self) -> str:
        return "".join(self._chunks)


@dataclass(slots=True)
class SheetLayout:
    header_row: int
    column_map: dict[str, int]


@dataclass(slots=True)
class SubmissionRecord:
    response: str
    submitted_at: datetime | None
    source_name: str


class ImportService:
    def inspect_workbook(self, workbook_path: Path) -> WorkbookProfile:
        course_name = workbook_path.stem.replace("_", " ").strip() or "Course Workbook"
        if not workbook_path.exists():
            return WorkbookProfile(
                path=workbook_path,
                course_name=course_name,
                assignments=[],
                semester_suggestions=self._default_semester_suggestions([]),
            )

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        assignments: list[AssignmentMeta] = []
        semesters: set[str] = set()

        for sheet in workbook.worksheets:
            if sheet.title == ADVANCED_HISTORY_SHEET:
                continue
            header_row = self._detect_header_row(sheet)
            prompt = self._existing_prompt(sheet)
            assignment_name = self._existing_assignment_name(sheet)
            prompt_summary = self._prompt_summary(assignment_name or prompt or sheet.title)
            existing_semesters = self._extract_semesters(sheet, header_row)
            response_count = self._count_responses(sheet, header_row)
            semesters.update(existing_semesters)
            assignments.append(
                AssignmentMeta(
                    sheet_name=sheet.title,
                    match_title=self._existing_match_title(sheet),
                    prompt=prompt,
                    prompt_summary=prompt_summary,
                    existing_semesters=existing_semesters,
                    response_count=response_count,
                )
            )

        workbook.close()
        return WorkbookProfile(
            path=workbook_path,
            course_name=course_name,
            assignments=assignments,
            semester_suggestions=self._default_semester_suggestions(sorted(semesters)),
        )

    def preview_archive(
        self,
        zip_path: Path,
        workbook_profile: WorkbookProfile | None = None,
        default_assignment: str = "",
        default_semester: str = "",
    ) -> ZipImportPreview:
        submissions = self.extract_submissions_from_zip(zip_path)
        response_count = len(submissions)
        submission_date = submissions[0].submitted_at if submissions else None
        detected_semester = self._semester_from_date(submission_date)
        if not detected_semester and default_semester.strip():
            detected_semester = default_semester.strip()
        detected = self.suggest_assignment(zip_path, workbook_profile)
        if not detected and default_assignment.strip():
            detected = default_assignment.strip()

        note = "Ready"
        if zip_path.suffix.lower() != ".zip":
            note = "Not a ZIP"
        elif not submissions:
            note = "No submission files detected"

        return ZipImportPreview(
            path=zip_path,
            response_count=response_count,
            detected_assignment=detected,
            detected_semester=detected_semester,
            submission_date=submission_date,
            note=note,
        )

    def create_course_workbook(
        self,
        workbook_path: Path,
        course: CourseCatalogItem,
    ) -> WorkbookProfile:
        workbook = Workbook()
        if "Sheet" in workbook.sheetnames:
            workbook.remove(workbook["Sheet"])

        assignment_names = course.assignment_names or ["Assignment 1"]
        for sheet_name in assignment_names:
            worksheet = workbook.create_sheet(self._safe_sheet_title(sheet_name))
            self._ensure_sheet_layout(worksheet, sheet_name)

        self._ensure_advanced_history_sheet(workbook)
        workbook_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(workbook_path)
        workbook.close()
        return self.inspect_workbook(workbook_path)

    def import_archives(
        self,
        workbook_path: Path,
        requests: list[ArchiveImportRequest],
        workbook_profile: WorkbookProfile | None = None,
        course: CourseCatalogItem | None = None,
        create_missing_sheets: bool = True,
        record_advanced_metadata: bool = True,
    ) -> list[ImportResult]:
        if not requests:
            raise ValueError("No ZIP archives were queued for import.")

        if not workbook_path.exists() and course:
            workbook_profile = self.create_course_workbook(workbook_path, course)

        workbook = self._load_or_create_workbook(workbook_path)
        self._ensure_advanced_history_sheet(workbook)
        batch_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        imported_at = datetime.now()
        results: list[ImportResult] = []

        for request in requests:
            target_sheet = self._safe_sheet_title(request.target_sheet.strip())
            if not target_sheet:
                raise ValueError(f"{request.zip_path.name}: target worksheet is empty.")
            semester = request.semester.strip()
            if not semester:
                raise ValueError(f"{request.zip_path.name}: semester is required.")

            submissions = self.extract_submissions_from_zip(request.zip_path)
            if not submissions:
                raise ValueError(f"{request.zip_path.name}: no readable submission files found.")

            worksheet, created_sheet = self._get_or_create_sheet(
                workbook,
                target_sheet,
                create_missing_sheets,
            )
            assignment_name, prompt_text, match_title = self._assignment_metadata_from_zip(
                request.zip_path,
                target_sheet,
            )
            layout = self._ensure_sheet_layout(worksheet, target_sheet)
            self._update_sheet_assignment_metadata(
                worksheet,
                assignment_name=assignment_name,
                match_title=match_title,
                prompt_text=prompt_text,
            )
            next_row = self._find_next_row(worksheet, layout.header_row)
            first_response_row = next_row

            for submission in submissions:
                worksheet.cell(next_row, layout.column_map["Semester"]).value = semester
                worksheet.cell(next_row, layout.column_map["Response"]).value = submission.response
                worksheet.cell(next_row, layout.column_map["Response"]).alignment = Alignment(
                    wrap_text=True,
                    vertical="top",
                )
                next_row += 1

            last_response_row = next_row - 1
            self._update_sheet_metadata(
                worksheet=worksheet,
                last_imported_line=last_response_row,
                imported_at=imported_at,
                semester=semester,
            )
            if record_advanced_metadata:
                self._append_advanced_history(
                    workbook=workbook,
                    batch_id=batch_id,
                    imported_at=imported_at,
                    semester=semester,
                    target_sheet=target_sheet,
                    source_zip=request.zip_path.name,
                    imported_count=len(submissions),
                    start_row=first_response_row,
                    end_row=last_response_row,
                )

            results.append(
                ImportResult(
                    zip_path=request.zip_path,
                    target_sheet=target_sheet,
                    imported_count=len(submissions),
                    created_sheet=created_sheet,
                    batch_id=batch_id,
                )
            )

        workbook_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(workbook_path)
        workbook.close()
        return results

    def extract_submissions_from_zip(self, zip_path: Path) -> list[SubmissionRecord]:
        if zip_path.suffix.lower() != ".zip" or not zip_path.exists():
            return []

        submissions: list[SubmissionRecord] = []
        with zipfile.ZipFile(zip_path) as archive:
            for name in archive.namelist():
                if name.endswith("/") or name.startswith("__MACOSX/"):
                    continue
                if not self._looks_like_submission(name):
                    continue

                raw = archive.read(name)
                text = self._decode_bytes(raw)
                if name.lower().endswith((".html", ".htm")):
                    text = self._strip_html(text)
                cleaned = self._normalize_response(text)
                if cleaned:
                    submissions.append(
                        SubmissionRecord(
                            response=cleaned,
                            submitted_at=self._submission_date_from_archive_name(name),
                            source_name=name,
                        )
                    )

        return submissions

    def suggest_assignment(
        self,
        zip_path: Path,
        workbook_profile: WorkbookProfile | None,
    ) -> str | None:
        if not workbook_profile or not workbook_profile.assignments:
            return None

        name = zip_path.stem
        number_match = re.search(r"(assignment|question|journal)\s*[-_ ]*(\d{1,2})", name, re.I)
        if number_match:
            target_number = number_match.group(2)
            for assignment in workbook_profile.assignments:
                match_source = assignment.match_title or assignment.prompt_summary or assignment.prompt or assignment.sheet_name
                source_match = re.search(r"(\d{1,2})", match_source)
                if source_match and source_match.group(1) == target_number:
                    return assignment.sheet_name

        best_score = 0.0
        best_sheet: str | None = None
        normalized_name = self._normalize_search_text(name)
        zip_keywords = self._keyword_set(name)

        for assignment in workbook_profile.assignments:
            match_title = assignment.match_title or assignment.prompt_summary or assignment.prompt or assignment.sheet_name
            haystack = f"{match_title} {assignment.prompt_summary}"
            normalized_sheet = self._normalize_search_text(haystack)
            overlap = len(zip_keywords & self._keyword_set(haystack))
            ratio = SequenceMatcher(None, normalized_name, normalized_sheet).ratio()
            score = ratio + (overlap * 0.1)

            if score > best_score:
                best_score = score
                best_sheet = assignment.sheet_name

        if best_score >= 0.35:
            return best_sheet
        return None

    def _load_or_create_workbook(self, workbook_path: Path) -> Workbook:
        if workbook_path.exists():
            return load_workbook(workbook_path)

        workbook = Workbook()
        if "Sheet" in workbook.sheetnames:
            default_sheet = workbook["Sheet"]
            if default_sheet.max_row == 1 and default_sheet.max_column == 1 and not default_sheet["A1"].value:
                workbook.remove(default_sheet)
        return workbook

    def _get_or_create_sheet(
        self,
        workbook: Workbook,
        sheet_name: str,
        create_missing_sheets: bool,
    ):
        if sheet_name in workbook.sheetnames:
            return workbook[sheet_name], False
        if not create_missing_sheets:
            raise ValueError(f"Worksheet '{sheet_name}' does not exist in the workbook.")
        return workbook.create_sheet(sheet_name), True

    def _ensure_sheet_layout(self, worksheet, assignment_name: str) -> SheetLayout:
        for index, label in enumerate(METADATA_LABELS, start=1):
            worksheet.cell(2, index).value = label
            if worksheet.cell(1, index).value in (None, ""):
                worksheet.cell(1, index).value = METADATA_NULL

        existing_assignment_name = self._existing_assignment_name(worksheet)
        existing_match_title = self._existing_match_title(worksheet)
        existing_prompt = self._existing_prompt(worksheet)
        worksheet.cell(ASSIGNMENT_NAME_ROW, 1).value = existing_assignment_name or assignment_name or worksheet.title
        worksheet.cell(ASSIGNMENT_NAME_ROW, MATCH_TITLE_COLUMN).value = existing_match_title
        worksheet.cell(PROMPT_ROW, 1).value = existing_prompt or "Assignment Text"
        if worksheet.cell(RUBRIC_ROW, 1).value in (None, ""):
            worksheet.cell(RUBRIC_ROW, 1).value = "Assignment Ruberic"

        for index, header in enumerate(WORKSHEET_HEADERS, start=1):
            worksheet.cell(HEADER_ROW, index).value = header
        for index, descriptor in enumerate(WORKSHEET_HEADER_DESCRIPTORS, start=1):
            worksheet.cell(DESCRIPTOR_ROW, index).value = descriptor

        worksheet.column_dimensions["A"].width = 21
        worksheet.column_dimensions["B"].width = 88
        worksheet.column_dimensions["C"].width = 24
        worksheet.column_dimensions["D"].width = 24
        worksheet.column_dimensions["E"].width = 22
        worksheet.column_dimensions["F"].width = 18
        worksheet.column_dimensions["G"].width = 18
        worksheet.column_dimensions["H"].width = 48
        worksheet.column_dimensions["I"].width = 48
        worksheet.column_dimensions["J"].width = 20
        worksheet.row_dimensions[PROMPT_ROW].height = max(worksheet.row_dimensions[PROMPT_ROW].height or 0, 36)
        worksheet.row_dimensions[RUBRIC_ROW].height = max(worksheet.row_dimensions[RUBRIC_ROW].height or 0, 36)
        worksheet.row_dimensions[HEADER_ROW].height = max(worksheet.row_dimensions[HEADER_ROW].height or 0, 20)
        worksheet.row_dimensions[DESCRIPTOR_ROW].height = max(worksheet.row_dimensions[DESCRIPTOR_ROW].height or 0, 42)

        worksheet.cell(PROMPT_ROW, 1).alignment = Alignment(wrap_text=True, vertical="top")
        worksheet.cell(RUBRIC_ROW, 1).alignment = Alignment(wrap_text=True, vertical="top")
        for column in range(1, len(WORKSHEET_HEADERS) + 1):
            worksheet.cell(DESCRIPTOR_ROW, column).alignment = Alignment(wrap_text=True, vertical="top")

        column_map = self._header_column_map(worksheet, HEADER_ROW)
        return SheetLayout(header_row=HEADER_ROW, column_map=column_map)

    def _update_sheet_assignment_metadata(
        self,
        worksheet,
        assignment_name: str,
        match_title: str,
        prompt_text: str,
    ) -> None:
        if assignment_name:
            worksheet.cell(ASSIGNMENT_NAME_ROW, 1).value = assignment_name
        if match_title:
            worksheet.cell(ASSIGNMENT_NAME_ROW, MATCH_TITLE_COLUMN).value = match_title
        if prompt_text:
            worksheet.cell(PROMPT_ROW, 1).value = prompt_text

    def _find_next_row(self, worksheet, header_row: int) -> int:
        last_row = max(worksheet.max_row, FIRST_DATA_ROW - 1)
        while last_row > header_row:
            has_value = any(
                worksheet.cell(last_row, column).value not in (None, "")
                for column in range(1, max(worksheet.max_column, len(WORKSHEET_HEADERS)) + 1)
            )
            if has_value:
                break
            last_row -= 1
        return max(last_row + 1, FIRST_DATA_ROW)

    def _extract_semesters(self, worksheet, header_row: int) -> list[str]:
        semester_values: set[str] = set()
        for row in range(header_row + 1, worksheet.max_row + 1):
            value = self._clean_text(worksheet.cell(row, 1).value)
            if not value:
                continue
            semester_values.add(value)
        return self._sort_semesters(semester_values)

    def _count_responses(self, worksheet, header_row: int) -> int:
        count = 0
        for row in range(header_row + 1, worksheet.max_row + 1):
            semester = self._clean_text(worksheet.cell(row, 1).value)
            response = self._clean_text(worksheet.cell(row, 2).value)
            if semester and response:
                count += 1
        return count

    def _detect_header_row(self, worksheet) -> int:
        if self._clean_text(worksheet.cell(HEADER_ROW, 1).value).casefold() == "semester":
            return HEADER_ROW
        for row_index in range(1, min(20, worksheet.max_row) + 1):
            values = [
                self._clean_text(worksheet.cell(row_index, column).value).casefold()
                for column in range(1, min(worksheet.max_column, 12) + 1)
            ]
            if "semester" in values and "response" in values:
                return row_index
        return HEADER_ROW

    def _header_column_map(self, worksheet, header_row: int) -> dict[str, int]:
        column_map: dict[str, int] = {}
        for column in range(1, worksheet.max_column + 1):
            value = self._clean_text(worksheet.cell(header_row, column).value)
            if value:
                column_map[value] = column
        for index, header in enumerate(WORKSHEET_HEADERS, start=1):
            column_map.setdefault(header, index)
        return column_map

    def _ensure_advanced_history_sheet(self, workbook: Workbook):
        if ADVANCED_HISTORY_SHEET in workbook.sheetnames:
            worksheet = workbook[ADVANCED_HISTORY_SHEET]
        else:
            worksheet = workbook.create_sheet(ADVANCED_HISTORY_SHEET)

        headers = (
            "Imported At",
            "Import Batch",
            "Semester",
            "Assignment",
            "Source ZIP",
            "Imported Count",
            "Start Row",
            "End Row",
        )
        for index, header in enumerate(headers, start=1):
            worksheet.cell(1, index).value = header
        return worksheet

    def _append_advanced_history(
        self,
        workbook: Workbook,
        batch_id: str,
        imported_at: datetime,
        semester: str,
        target_sheet: str,
        source_zip: str,
        imported_count: int,
        start_row: int,
        end_row: int,
    ) -> None:
        worksheet = self._ensure_advanced_history_sheet(workbook)
        row_index = worksheet.max_row + 1
        values = (
            imported_at.strftime("%Y-%m-%d %H:%M:%S"),
            batch_id,
            semester,
            target_sheet,
            source_zip,
            imported_count,
            start_row,
            end_row,
        )
        for index, value in enumerate(values, start=1):
            worksheet.cell(row_index, index).value = value

    def _update_sheet_metadata(
        self,
        worksheet,
        last_imported_line: int,
        imported_at: datetime,
        semester: str,
    ) -> None:
        worksheet.cell(1, 1).value = last_imported_line
        worksheet.cell(1, 3).value = imported_at.date()
        worksheet.cell(1, 5).value = semester
        for column in (2, 4, 6):
            if worksheet.cell(1, column).value in (None, ""):
                worksheet.cell(1, column).value = METADATA_NULL

    def _looks_like_submission(self, archive_name: str) -> bool:
        lower_name = archive_name.lower()
        base_name = Path(lower_name).name
        if base_name in {"submission.txt", "submission.html", "submission.htm"}:
            return True
        if base_name.startswith("submission.") and base_name.endswith((".txt", ".html", ".htm")):
            return True
        if "submission" in base_name and base_name.endswith((".txt", ".html", ".htm")):
            return True
        return False

    def _decode_bytes(self, raw_bytes: bytes) -> str:
        encodings = ("utf-8", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "latin-1")
        for encoding in encodings:
            try:
                return raw_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw_bytes.decode("utf-8", errors="ignore")

    def _strip_html(self, raw_text: str) -> str:
        stripped = _HTMLStripper()
        stripped.feed(raw_text)
        return html.unescape(stripped.get_text())

    def _normalize_response(self, text: str) -> str:
        lines = [line.strip() for line in text.replace("\r", "\n").split("\n")]
        collapsed = "\n".join(line for line in lines if line)
        return collapsed.strip()

    def _clean_text(self, value) -> str:
        if value is None:
            return ""
        return str(value).replace("\r", " ").replace("\n", " ").strip()

    def _existing_assignment_name(self, worksheet) -> str:
        values = (
            self._clean_text(worksheet.cell(ASSIGNMENT_NAME_ROW, 1).value),
            self._clean_text(worksheet.cell(4, 1).value),
        )
        for value in values:
            if value and value.casefold() != "history line (save for grading history attempts)":
                return value
        return ""

    def _existing_match_title(self, worksheet) -> str:
        values = (
            self._clean_text(worksheet.cell(ASSIGNMENT_NAME_ROW, MATCH_TITLE_COLUMN).value),
            self._clean_text(worksheet.cell(4, MATCH_TITLE_COLUMN).value),
        )
        for value in values:
            if value and value.casefold() != "assignment text":
                return value
        return ""

    def _existing_prompt(self, worksheet) -> str:
        values = (
            self._clean_text(worksheet.cell(PROMPT_ROW, 1).value),
            self._clean_text(worksheet.cell(5, 1).value),
        )
        for value in values:
            if value and value.casefold() != "assignment text":
                return value
        return ""

    def _prompt_summary(self, prompt: str, limit: int = 82) -> str:
        prompt = self._clean_text(prompt)
        if len(prompt) <= limit:
            return prompt
        return prompt[: limit - 3].rstrip() + "..."

    def _assignment_metadata_from_zip(self, zip_path: Path, target_sheet: str) -> tuple[str, str, str]:
        raw_name = zip_path.stem
        raw_name = re.sub(
            r"\s+Download\s+[A-Z][a-z]{2,8}\s+\d{1,2},\s+\d{4}\s+\d{1,4}\s*[AP]M$",
            "",
            raw_name,
            flags=re.I,
        ).strip()
        match_title = re.sub(r"\s+", " ", raw_name).strip(" -_:")
        cleaned = raw_name
        for phrase in GENERIC_ASSIGNMENT_PHRASES:
            cleaned = re.sub(rf"\b{re.escape(phrase)}\b", " ", cleaned, flags=re.I)
        cleaned = re.sub(r"[_\-]+", " ", cleaned)
        cleaned = re.sub(r"\s+", " ", cleaned).strip(" -_:")

        prompt_text = cleaned or target_sheet
        if not prompt_text:
            prompt_text = "Assignment Text"

        if prompt_text and prompt_text.casefold() not in target_sheet.casefold():
            assignment_name = f"{target_sheet} - {prompt_text}"
        else:
            assignment_name = target_sheet or prompt_text
        assignment_name = re.sub(r"\s+", " ", assignment_name).strip(" -_:")
        return assignment_name or target_sheet, prompt_text, match_title or prompt_text

    def _normalize_search_text(self, value: str) -> str:
        return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()

    def _keyword_set(self, value: str) -> set[str]:
        return {
            token
            for token in self._normalize_search_text(value).split()
            if len(token) > 2 and token not in STOP_WORDS
        }

    def _default_semester_suggestions(self, workbook_semesters: list[str]) -> list[str]:
        suggestions = self._sort_semesters(workbook_semesters)
        now = datetime.now()
        year = now.year
        auto_suggestions = [
            self._format_semester_label("spring", year),
            self._format_semester_label("summer", year),
            self._format_semester_label("fall", year),
            self._format_semester_label("spring", year + 1),
        ]
        for value in auto_suggestions:
            if value not in suggestions:
                suggestions.append(value)
        return suggestions

    def _sort_semesters(self, semester_values) -> list[str]:
        def sort_key(value: str):
            match = re.search(r"(spring|summer|fall|winter)\s+(\d{2,4})", value, re.I)
            if not match:
                return (-1, value.casefold())
            season = match.group(1).casefold()
            year = int(match.group(2))
            if year < 100:
                year += 2000
            return (year, SEASON_ORDER.get(season, -1))

        return sorted({value for value in semester_values if value}, key=sort_key)

    def _safe_sheet_title(self, value: str) -> str:
        title = re.sub(r"[\[\]:*?/\\]", "-", value).strip()
        if len(title) > 31:
            title = title[:31].rstrip()
        return title

    def _submission_date_from_archive_name(self, archive_name: str) -> datetime | None:
        parts = [part.strip() for part in archive_name.split("/") if part.strip()]
        if len(parts) < 2:
            return None

        folder_name = parts[-2]
        match = re.search(
            r" - "
            r"([A-Z][a-z]{2,8}\s+\d{1,2},\s+\d{4}\s+\d{1,4}\s*[AP]M)$",
            folder_name,
        )
        if not match:
            return None

        raw_date = match.group(1)
        normalized = re.sub(r"\s+", " ", raw_date).strip()
        for fmt in ("%b %d, %Y %I%M %p", "%B %d, %Y %I%M %p"):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                continue
        return None

    def _semester_from_date(self, value: datetime | None) -> str | None:
        if value is None:
            return None

        if value.month in (1, 2, 3, 4, 5):
            season = "spring"
        elif value.month in (6, 7, 8):
            season = "summer"
        elif value.month in (9, 10, 11, 12):
            season = "fall"
        else:
            return None
        return self._format_semester_label(season, value.year)

    def _format_semester_label(self, season: str, year: int) -> str:
        return f"{season.title()} {year % 100:02d}"
